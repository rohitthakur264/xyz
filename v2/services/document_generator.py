import os
import json
import tempfile
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from database.schema_loader import SchemaLoader
from database.schema_anchors import find_anchor_tables, PATIENT_FK_CANDIDATES, get_fk_join
from database.query_executor import QueryExecutor
from database.connection_pool import get_connection
from config.logging import get_logger
from config.settings import AUDIT_PREFIXES
from datetime import datetime

logger = get_logger(__name__)

NORMAL_RANGES = {}
try:
    nr_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "config", "normal_ranges.json")
    if os.path.exists(nr_path):
        with open(nr_path, "r") as f:
            NORMAL_RANGES = json.load(f)
except Exception as e:
    logger.warning(f"Could not load normal ranges: {e}")

# Load value mappings for coded fields (gender, marital status, etc.)
_VALUE_MAPPINGS = {}
try:
    _vm_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "config", "value_mappings.json")
    if os.path.exists(_vm_path):
        with open(_vm_path, "r") as f:
            _VALUE_MAPPINGS = json.load(f)
except Exception as e:
    logger.warning(f"Could not load value_mappings.json: {e}")

REPORT_SCHEMA = {}
try:
    _rs_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "config", "report_schema.json")
    if os.path.exists(_rs_path):
        with open(_rs_path, "r") as f:
            REPORT_SCHEMA = json.load(f)
except Exception as e:
    logger.warning(f"Could not load report_schema.json: {e}")

def validate_report_schema(schema: dict, anchors: dict):
    """Validates the report schema against the live schema cache at startup."""
    import re
    if not REPORT_SCHEMA or "sections" not in REPORT_SCHEMA:
        logger.warning("No report sections found to validate.")
        return
        
    for sec in REPORT_SCHEMA["sections"]:
        anchor_key = sec["anchor"]
        if anchor_key not in anchors:
            raise ValueError(f"Report section '{sec['id']}' relies on anchor '{anchor_key}' which is not in table_mappings.json")
            
        main_table = anchors[anchor_key]
        if isinstance(main_table, dict):
            main_table = main_table.get("table")
        if main_table not in schema:
            raise ValueError(f"Report section '{sec['id']}' anchor table '{main_table}' not in live schema.")
            
        aliases = {sec["alias"]: main_table}
        
        # Validate joins and their FKs
        for j in sec.get("joins", []):
            j_anchor = j["anchor"]
            if j_anchor not in anchors:
                raise ValueError(f"Join anchor '{j_anchor}' not found in table_mappings.json.")
            j_table = anchors[j_anchor]
            if isinstance(j_table, dict):
                j_table = j_table.get("table")
            if j_table not in schema:
                raise ValueError(f"Join table '{j_table}' not in live schema.")
                
            aliases[j["alias"]] = j_table
            from_alias = j.get("from_alias", sec["alias"])
            if from_alias not in aliases:
                raise ValueError(f"Join from_alias '{from_alias}' not defined before use.")
            
            from_table = aliases[from_alias]
            via = j.get("via_column")
            try:
                get_fk_join(from_table, j_table, schema, via_column=via)
            except Exception as e:
                raise ValueError(f"Report section '{sec['id']}': {e}")
                
        # Validate columns
        col_exprs = []
        for c in sec.get("select", []) + sec.get("unscoped_select", []):
            if "column" in c:
                col_exprs.append(c["column"])
            elif "display_columns" in c:
                col_exprs.extend(c["display_columns"])
        if sec.get("where_extra"):
            col_exprs.append(sec["where_extra"])
        if sec.get("order_by"):
            col_exprs.append(sec["order_by"])
            
        for expr in col_exprs:
            matches = re.findall(r'\b([a-zA-Z0-9_]+)\.([a-zA-Z0-9_]+)\b', expr)
            for alias, col_name in matches:
                if alias in aliases:
                    t_name = aliases[alias]
                    t_cols = [c["column_name"].lower() for c in schema.get(t_name, {}).get("columns", [])]
                    if col_name.lower() not in t_cols:
                        raise ValueError(f"Report section '{sec['id']}': column '{col_name}' not found on table '{t_name}' (alias '{alias}').")
                        
    logger.info("Validated report schema successfully.")

class DocumentGenerator:
    """Generates structured medical records as structured dict, Markdown, or Word."""
    
    def __init__(self, query_executor=None, anchors=None, schema=None):
        from database.schema_anchors import load_table_mappings, merge_extra_relationships
        self.executor = query_executor or QueryExecutor()
        if schema is not None:
            self.schema = schema
        else:
            self.schema = SchemaLoader().load_schema()
            merge_extra_relationships(self.schema)
        self.anchors = anchors if anchors is not None else load_table_mappings(self.schema)
        self.patient_table = self._require_table("patient")
        self.visit_table = self._require_table("visit")

    def _require_table(self, key: str) -> str:
        """Return a table configured for an anchor. No hospital-specific fallback is used."""
        entry = self.anchors.get(key)
        if isinstance(entry, dict):
            table = entry.get("table")
        else:
            table = entry
        if not table or not str(table).strip():
            raise ValueError(f"Missing required schema anchor: {key}")
        return str(table).strip()

    def _anchor_config(self, key: str) -> dict:
        """Return normalized anchor configuration."""
        entry = self.anchors.get(key)
        if isinstance(entry, dict):
            table = entry.get("table")
            if not table:
                raise ValueError(f"Missing table in schema anchor: {key}")
            return entry
        if entry:
            cfg = {"table": str(entry).strip()}
            self.anchors[key] = cfg
            return cfg
        raise ValueError(f"Missing required schema anchor: {key}")

    def _lookup_display(self, anchor, id_val) -> str:
        """Resolve a coded value using anchor metadata or generic display/id columns."""
        if not id_val or str(id_val).strip() in ("0", "-1", ""):
            return "—"

        if isinstance(anchor, str):
            cfg = {"table": anchor}
        elif isinstance(anchor, dict):
            cfg = anchor
        else:
            return "—"

        table = cfg.get("table")
        if not table:
            return "—"

        try:
            id_col = cfg.get("id_col")
            display_col = cfg.get("display_col")

            if not id_col or not display_col:
                conn = get_connection()
                cur = conn.cursor(dictionary=True)
                cur.execute(
                    """
                    SELECT COLUMN_NAME, DATA_TYPE
                    FROM INFORMATION_SCHEMA.COLUMNS
                    WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s
                    ORDER BY ORDINAL_POSITION
                    """,
                    (table,)
                )
                columns = cur.fetchall()
                cur.close()
                conn.close()

                names = [c["COLUMN_NAME"] for c in columns]
                lower = {n.lower(): n for n in names}

                if not id_col:
                    id_col = next(
                        (lower[n] for n in ("id", "code", "serviceid", "empid", "departmentid")
                         if n in lower),
                        None
                    )

                if not display_col:
                    display_col = next(
                        (lower[n] for n in ("display", "name", "description", "servicename")
                         if n in lower),
                        None
                    )

                if not id_col or not display_col:
                    logger.warning(
                        f"Could not resolve lookup columns for anchor '{table}'"
                    )
                    return "—"
                
                # CACHE THE RESOLVED COLUMNS to prevent hitting INFORMATION_SCHEMA again!
                cfg["id_col"] = id_col
                cfg["display_col"] = display_col

            res = self.executor.execute_prepared(
                f"SELECT `{display_col}` AS display FROM `{table}` "
                f"WHERE `{id_col}` = %s LIMIT 1",
                (id_val,)
            )
            if res.get("success") and res.get("rows"):
                return str(res["rows"][0].get("display", "—"))

        except Exception as e:
            logger.debug(f"Lookup failed for {table} id={id_val}: {e}")

        return "—"

    def _clean_raw_data(self, rows: list, is_vitals=False) -> list:
        if not rows:
            return []
        cleaned_rows = []
        audit_prefixes = tuple(p.strip().lower() for p in AUDIT_PREFIXES if p.strip())
        for row in rows:
            if isinstance(row, dict):
                clean_r = {}
                for k, v in row.items():
                    if v is None or v == "" or str(v).lower() in ("null", "none"):
                        continue
                    if any(k.lower().startswith(p) or k.lower().endswith(p) for p in audit_prefixes):
                        continue
                    clean_r[k] = v
                    
                    if is_vitals:
                        for nr_key, nr_val in NORMAL_RANGES.items():
                            if nr_key in k.lower():
                                clean_r[f"{k} Normal Range"] = f"{nr_val['range']} {nr_val['unit']}"
                cleaned_rows.append(clean_r)
        return cleaned_rows

    def _map_gender(self, gender_code) -> str:
        """Map gender code to label using DB lookup, with config fallback."""
        if not gender_code or str(gender_code).strip() in ("0", ""):
            return "—"
        # Try DB lookup first
        gender_table = self.anchors.get("gender")
        if gender_table:
            label = self._lookup_display(self._anchor_config("gender"), gender_code)
            if label != "—":
                return label
        # Config fallback
        gender_map = _VALUE_MAPPINGS.get("gender", {})
        return gender_map.get(str(gender_code), "—")

    def _fetch_rich_demographics(self, patient_id: int) -> dict:
        """Fetches full patient demographics using report_schema and applies shaping logic."""
        try:
            sec_configs = [s for s in REPORT_SCHEMA.get("sections", []) if s.get("id", "").startswith("demographics")]
            if not sec_configs:
                return {}
                
            row = {}
            for sec_config in sec_configs:
                sql, params = self._build_section_query(sec_config, patient_id)
                res = self.executor.execute_prepared(sql, params)
                if not res.get("success") or not res.get("rows"):
                    continue
                    
                cleaned_rows = self._clean_raw_data(res.get("rows"))
                r = cleaned_rows[0] if cleaned_rows else res.get("rows")[0]
                
                if r:
                    row.update({k.lower(): v for k, v in r.items()})
            
            if not row:
                return {}
            
            # Build clean demographics dict
            full_name_parts = [
                row.get("first_name", ""),
                row.get("middle_name", ""),
                row.get("thirdname", ""),
                row.get("last_name", "")
            ]
            full_name = " ".join(p for p in full_name_parts if p and str(p).strip())
            
            # Arabic name
            ar_parts = [
                row.get("r_firstname", ""),
                row.get("r_middlename", ""),
                row.get("r_thirdname", ""),
                row.get("r_familyname", "")
            ]
            ar_name = " ".join(p for p in ar_parts if p and str(p).strip())
            
            dob = row.get("date_of_birth")
            dob_str = dob.strftime("%d-%b-%Y") if dob else "—"
            
            age_str = str(row.get("age", "—"))
            if row.get("ageinmonths"):
                age_str += f" yrs {row['ageinmonths']} months"
            else:
                age_str += " yrs"
            
            demographics = {
                "Full Name (English)": full_name or row.get("display", "—"),
                "Full Name (Arabic)": ar_name or "—",
                "Registration No": str(row.get("registration_no", "—")),
                "Patient ID": str(row.get("patientid", "—")),
                "Gender": row.get("gender_display") or self._map_gender(row.get("gender")),
                "Date of Birth": dob_str,
                "Age": age_str,
                "Nationality": row.get("nationality_display") or str(row.get("nationality", "—")),
                "Marital Status": row.get("marital_status_display") or str(row.get("maritalstatus", "—")),
                "Mobile No.": str(row.get("mobile", "—")),
                "Alternate Mobile": str(row.get("alternate_mobile", "—")) if row.get("alternate_mobile") else "—",
                "Passport No.": str(row.get("passportno", "—")) if row.get("passportno") else "—",
                "National No.": str(row.get("nationalno", "—")) if row.get("nationalno") else "—",
                "Email": str(row.get("emailaddress", "—")) if row.get("emailaddress") else "—",
            }
            
            # Next of Kin section
            nok_name_parts = [
                row.get("nokfirstname", ""),
                row.get("nokmiddlename", ""),
                row.get("noklastname", ""),
                row.get("nokfamilyname", "")
            ]
            nok_name = " ".join(p for p in nok_name_parts if p and str(p).strip())
            
            nok = {
                "Next of Kin Name": nok_name or "—",
                "NOK Relationship": str(row.get("nokrelationship", "—")) if row.get("nokrelationship") else "—",
                "NOK Mobile": str(row.get("nokmobile", "—")) if row.get("nokmobile") else "—",
                "NOK Email": str(row.get("nokemail", "—")) if row.get("nokemail") else "—",
            }
            
            # Remove dash-only entries
            demographics = {k: v for k, v in demographics.items() if v and v != "—"}
            nok = {k: v for k, v in nok.items() if v and v != "—"}
            
            return {"demographics": demographics, "next_of_kin": nok}
            
        except Exception as e:
            logger.error(f"Failed to fetch rich demographics: {e}", exc_info=True)
            return {}

    def fetch_visit_info(self, patient_id: int) -> list:
        """Fetches visit details with department and doctor lookups."""
        try:
            sec_config = next((s for s in REPORT_SCHEMA.get("sections", []) if s.get("id") == "visits"), None)
            if not sec_config:
                return []
                
            sql, params = self._build_section_query(sec_config, patient_id)
            res = self.executor.execute_prepared(sql, params)
            if not res.get("success") or not res.get("rows"):
                return []
                
            visits = self._clean_raw_data(res.get("rows"))
            
            result = []
            for v in visits:
                v_lower = {k.lower(): val for k, val in v.items()}
                visit_info = {
                    "Visit ID": str(v_lower.get("visit_id", "—")),
                    "Visit No": str(v_lower.get("visitno", "—")),
                    "Visit Type": str(v_lower.get("typeofvisit", "—")),
                    "Visit Date/Time": v_lower["visitdate"].strftime("%d-%b-%Y %H:%M") if v_lower.get("visitdate") else "—",
                    "Admission No.": "—", # Handled elsewhere if admission exists, but we'll default to — for OPD
                    "Department": v_lower.get("department_name") or (f"Department ID: {v_lower.get('departmentid', '')}" if v_lower.get("departmentid") else "—"),
                    "Attending Doctor": v_lower.get("doctor_name") or (f"Doctor ID: {v_lower.get('docincharge', '')}" if v_lower.get("docincharge") else "—"),
                    "Ward / Room / Bed": "—",
                    "Age at Visit": str(v_lower.get("age", "—")),
                    "Follow-up (days)": str(v_lower.get("followupdays", "—")) if v_lower.get("followupdays") else "—",
                }
                
                # Ward/Room/Bed — only include if non-zero (OPD patients have 0)
                if v_lower.get("wardid") and int(v_lower.get("wardid", 0)) > 0:
                    ward = self._lookup_display(self.anchors.get("ward"), v_lower["wardid"])
                    room = self._lookup_display(self.anchors.get("room"), v_lower.get("roomid"))
                    bed = self._lookup_display(self.anchors.get("bed"), v_lower.get("bedid"))
                    
                    if ward == "—":
                        ward = str(v_lower.get("wardid", "—"))
                        
                    visit_info["Ward / Room / Bed"] = f"{ward} / {room} / {bed}"
                
                result.append(visit_info)
            
            return result
            
        except Exception as e:
            logger.error(f"Failed to fetch visit info: {e}", exc_info=True)
            return []

    def _fetch_clinical_assessment(self, patient_id: int) -> list:
        """Fetches clinical assessments as key-value records."""
        try:
            sec_config = next((s for s in REPORT_SCHEMA.get("sections", []) if s.get("id") == "clinical_assessments"), None)
            if not sec_config:
                return []
                
            sql, params = self._build_section_query(sec_config, patient_id)
            res = self.executor.execute_prepared(sql, params)
            if not res.get("success") or not res.get("rows"):
                return []
            
            rows = self._clean_raw_data(res.get("rows"))
            
            result = []
            for row in rows:
                row_lower = {k.lower(): val for k, val in row.items()}
                assessment = {}
                field_labels = {
                    "reasonforconsultation": "Reason for Consultation",
                    "historyofillness": "History of Present Illness",
                    "initialdignosis": "Initial Diagnosis (ICD)",
                    "treatmentplan": "Treatment Plan",
                    "significantsigns": "Significant Signs",
                    "generalexamination": "General Examination",
                    "pastmedicalsurgical": "Past Medical/Surgical History",
                    "durationofillness": "Duration of Illness",
                    "painscore": "Pain Score",
                    "painsite": "Pain Site",
                }
                for db_col, label in field_labels.items():
                    val = row_lower.get(db_col)
                    if val and str(val).strip() and str(val).strip() != "0":
                        assessment[label] = str(val).strip()
                
                if assessment:
                    result.append(assessment)
            
            return result
            
        except Exception as e:
            logger.error(f"Failed to fetch clinical assessment: {e}", exc_info=True)
            return []

    def _build_section_query(self, sec_config: dict, patient_id: int, target_visit_id: int = None) -> tuple[str, tuple]:
        anchor_key = sec_config["anchor"]
        main_table = self._require_table(anchor_key)
        main_alias = sec_config["alias"]
        aliases = {main_alias: main_table}
        
        # Build SELECT
        select_items = []
        for c in sec_config.get("select", []):
            if "column" in c:
                select_items.append(f"{c['column']} AS '{c['alias']}'")
            elif "display_columns" in c:
                compose_rule = c.get("compose", "concat_space")
                if compose_rule == "concat_space":
                    expr = f"CONCAT_WS(' ', {', '.join(c['display_columns'])})"
                elif compose_rule == "concat_comma":
                    expr = f"CONCAT_WS(', ', {', '.join(c['display_columns'])})"
                else:
                    expr = c["display_columns"][0]
                select_items.append(f"{expr} AS '{c['alias']}'")
            
        if not target_visit_id:
            for c in sec_config.get("unscoped_select", []):
                if "column" in c:
                    select_items.append(f"{c['column']} AS '{c['alias']}'")
        
        select_clause = ", ".join(select_items)
        
        # Build FROM and JOINS
        from_clause = f"FROM {main_table} {main_alias}"
        join_clauses = []
        
        for j in sec_config.get("joins", []):
            j_anchor = j["anchor"]
            if j_anchor not in self.anchors:
                continue
            j_table = self._require_table(j_anchor)
            j_alias = j["alias"]
            j_type = j.get("type", "LEFT JOIN")
            from_alias = j.get("from_alias", main_alias)
            from_table = aliases[from_alias]
            
            try:
                fk = get_fk_join(from_table, j_table, self.schema, via_column=j.get("via_column"))
                cast_type = j.get("cast_join")
                if cast_type:
                    join_clauses.append(f"{j_type} {j_table} {j_alias} ON CAST({from_alias}.{fk['column_name']} AS {cast_type}) = CAST({j_alias}.{fk['referenced_column']} AS {cast_type})")
                else:
                    join_clauses.append(f"{j_type} {j_table} {j_alias} ON {from_alias}.{fk['column_name']} = {j_alias}.{fk['referenced_column']}")
                aliases[j_alias] = j_table
            except Exception as e:
                logger.error(f"Failed to resolve join in report section {sec_config['id']}: {e}")
                
        if not target_visit_id:
            # If any unscoped column needs visit, join visit
            needs_visit = any(c.get("join_visit") for c in sec_config.get("unscoped_select", []))
            if needs_visit:
                try:
                    fk = get_fk_join(main_table, self.visit_table, self.schema)
                    join_clauses.append(f"LEFT JOIN {self.visit_table} v ON {main_alias}.{fk['column_name']} = v.{fk['referenced_column']}")
                except Exception:
                    pass
                    
        joins_str = " ".join(join_clauses)
        
        # Build WHERE
        try:
            p_fk = get_fk_join(main_table, self.patient_table, self.schema)
            where_clause = f"WHERE {main_alias}.{p_fk['column_name']} = %s"
            params = [patient_id]
        except Exception:
            if main_table == self.patient_table:
                p_pk = self.schema.get(main_table, {}).get("primary_key", "id")
                where_clause = f"WHERE {main_alias}.{p_pk} = %s"
                params = [patient_id]
            else:
                where_clause = f"WHERE {main_alias}.patientid = %s"
                params = [patient_id]
                
        if target_visit_id:
            if main_table == self.visit_table:
                v_pk = self.schema.get(main_table, {}).get("primary_key", "visit_id")
                where_clause += f" AND {main_alias}.{v_pk} = %s"
                params.append(target_visit_id)
            else:
                try:
                    v_fk = get_fk_join(main_table, self.visit_table, self.schema)
                    where_clause += f" AND {main_alias}.{v_fk['column_name']} = %s"
                    params.append(target_visit_id)
                except Exception:
                    where_clause += f" AND {main_alias}.visitid = %s"
                    params.append(target_visit_id)
                    
        extra_where = sec_config.get("where_extra")
        if extra_where:
            where_clause += f" AND {extra_where}"
            
        # Build ORDER BY
        order_by = sec_config.get("order_by")
        order_clause = f"ORDER BY {order_by}" if order_by else ""
        
        limit = sec_config.get("limit")
        limit_clause = f"LIMIT {limit}" if limit else ""
        
        sql = f"SELECT {select_clause} {from_clause} {joins_str} {where_clause} {order_clause} {limit_clause}"
        return sql.strip(), tuple(params)

    def build_record_data(self, patient_context, target_visit_id: int = None) -> dict:
        """Builds a comprehensive dictionary of patient record data."""
        patient_id = int(patient_context.patient_id)
        
        # Fetch rich demographics from DB with JOINs
        rich_demo = self._fetch_rich_demographics(patient_id)
        
        data = {
            "patient": {
                "name": patient_context.name,
                "registration_no": patient_context.registration_no,
                "patient_id": patient_context.patient_id,
                "demographics": rich_demo.get("demographics", {}),
                "next_of_kin": rich_demo.get("next_of_kin", {}),
            },
            "visits": self.fetch_visit_info(patient_id),
            "clinical_assessments": self._fetch_clinical_assessment(patient_id),
            "sections": []
        }
        
        sections = []
        if not REPORT_SCHEMA or "sections" not in REPORT_SCHEMA:
            logger.error("REPORT_SCHEMA is not loaded or missing 'sections'.")
            return data
            
        for sec_config in REPORT_SCHEMA["sections"]:
            anchor_key = sec_config["anchor"]
            if anchor_key not in self.anchors:
                continue
                
            if sec_config.get("id", "").startswith("demographics") or sec_config.get("id") in ["visits", "clinical_assessments"]:
                continue
                
            sql, params = self._build_section_query(sec_config, patient_id, target_visit_id)
            
            sections.append({
                "id": sec_config["id"],
                "title": sec_config["title"],
                "sql": sql,
                "params": params,
                "is_vitals": sec_config.get("is_vitals", False),
                "format": sec_config.get("format", "table")
            })
            
        import concurrent.futures
        
        def fetch_section(sec):
            res = self.executor.execute_prepared(sec["sql"], sec["params"])
            if not res.get("success") or not res.get("rows"):
                return {"title": sec["title"], "rows": [], "error": "No records found.", "id": sec["id"], "format": sec["format"]}
                
            rows = self._clean_raw_data(res.get("rows"), is_vitals=sec["is_vitals"])
            if not rows:
                return {"title": sec["title"], "rows": [], "error": "No valid records found.", "id": sec["id"], "format": sec["format"]}
                
            return {"title": sec["title"], "rows": rows, "error": None, "id": sec["id"], "format": sec["format"]}
            
        with concurrent.futures.ThreadPoolExecutor(max_workers=len(sections) if sections else 1) as pool:
            results = list(pool.map(fetch_section, sections))
            data["sections"].extend(results)
            
        return data

    def _render_kv_markdown(self, kv_dict: dict, indent: str = "") -> list:
        """Renders a key-value dict as markdown bold-key lines."""
        lines = []
        for k, v in kv_dict.items():
            val_str = str(v).replace("\n", "  \n").strip()
            lines.append(f"{indent}- **{k}:** {val_str}")
        return lines

    def _render_table_markdown(self, kv_dict: dict) -> list:
        """Renders a key-value dict as a 4-column markdown table."""
        lines = []
        lines.append("| Field | Value | Field | Value |")
        lines.append("|---|---|---|---|")
        
        items = list(kv_dict.items())
        for i in range(0, len(items), 2):
            k1, v1 = items[i]
            v1_str = str(v1).replace("\n", " ").strip()
            
            if i + 1 < len(items):
                k2, v2 = items[i+1]
                v2_str = str(v2).replace("\n", " ").strip()
            else:
                k2, v2_str = "", ""
                
            lines.append(f"| **{k1}** | {v1_str} | **{k2}** | {v2_str} |")
        return lines

    def _render_visit_html(self, data: dict) -> list:
        """Builds a precise 4-column HTML table matching the provided layout."""
        items = list(data.items())
        if len(items) % 2 != 0:
            items.append(("", ""))
            
        rows_html = []
        for i in range(0, len(items), 2):
            k1, v1 = items[i]
            k2, v2 = items[i+1]
            rows_html.append(
                f"<tr><th style='background:#f4f4f4; border:1px solid #ccc; padding:6px 10px; text-align:left; font-weight:bold;'>{k1}</th>"
                f"<td style='border:1px solid #ccc; padding:6px 10px;'>{v1}</td>"
                f"<th style='background:#f4f4f4; border:1px solid #ccc; padding:6px 10px; text-align:left; font-weight:bold;'>{k2}</th>"
                f"<td style='border:1px solid #ccc; padding:6px 10px;'>{v2}</td></tr>"
            )
            
        return [
            "<table style='width:100%; border-collapse:collapse; font-family:Arial, sans-serif; font-size:13px; margin-bottom:18px;'>",
            *rows_html,
            "</table>"
        ]

    def render_markdown(self, data: dict) -> str:
        """Renders the record data as a clean Markdown string."""
        lines = []
        lines.append("# PATIENT MEDICAL RECORD")
        lines.append("## Consolidated Clinical Summary")
        lines.append(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        lines.append("")
        
        # --- Demographics ---
        patient_name = str(data['patient'].get('name', '')).upper()
        lines.append(f"### 1. Patient: {patient_name}")
        lines.append("#### Demographics")
        demo = data["patient"].get("demographics", {})
        if demo:
            lines.extend(self._render_table_markdown(demo))
        else:
            lines.extend(self._render_table_markdown({
                "Registration No": data['patient']['registration_no'],
                "Patient ID": data['patient']['patient_id']
            }))
        lines.append("")
        
        # --- Next of Kin ---
        nok = data["patient"].get("next_of_kin", {})
        if nok:
            lines.append("#### Next of Kin")
            lines.extend(self._render_table_markdown(nok))
            lines.append("")
        
        # --- Visit Info ---
        visits = data.get("visits", [])
        section_num = 2
        if visits:
            lines.append(f"### {section_num}. Visit History ({len(visits)} visit{'s' if len(visits) > 1 else ''})")
            for i, v_info in enumerate(visits, 1):
                visit_no = v_info.get("Visit No", f"Visit {i}")
                lines.append(f"#### Visit: {visit_no}")
                lines.extend(self._render_visit_html(v_info))
                lines.append("")
            section_num += 1
        
        # --- Clinical Assessment ---
        assessments = data.get("clinical_assessments", [])
        if assessments:
            lines.append(f"### {section_num}. Clinical Assessment ({len(assessments)} record{'s' if len(assessments) > 1 else ''})")
            for i, asmt in enumerate(assessments, 1):
                if len(assessments) > 1:
                    lines.append(f"#### Assessment {i}")
                lines.extend(self._render_kv_markdown(asmt))
                lines.append("")
            section_num += 1
        
        # --- Tabular Sections (Vitals, Labs, Prescriptions) ---
        for sec in data["sections"]:
            row_count = len(sec['rows'])
            if row_count > 0:
                lines.append(f"### {section_num}. {sec['title']} ({row_count} readings/records)")
            else:
                lines.append(f"### {section_num}. {sec['title']}")
            section_num += 1
                
            if sec["error"]:
                lines.append(f"*{sec['error']}*")
                lines.append("")
                continue
                
            if sec.get("format") == "paragraph":
                for row_idx, row in enumerate(sec["rows"]):
                    if len(sec["rows"]) > 1:
                        lines.append(f"#### Record {row_idx + 1}")
                    
                    # Ensure nicely title-cased keys just in case
                    formatted_row = {str(k).replace('_', ' ').title(): v for k, v in row.items()}
                    lines.extend(self._render_kv_markdown(formatted_row))
                    lines.append("")
                continue

            # Create a markdown table
            headers = []
            for r in sec["rows"]:
                for k in r.keys():
                    if k not in headers:
                        headers.append(k)
            header_line = "| " + " | ".join([str(h).replace('_', ' ').title() for h in headers]) + " |"
            sep_line = "|" + "|".join(["---"] * len(headers)) + "|"
            
            lines.append(header_line)
            lines.append(sep_line)
            
            for row in sec["rows"]:
                row_vals = []
                
                # Check for critical/out of bounds flag if it's the Labs section
                is_critical_row = False
                if sec.get("id") == "labs":
                    crit_flag = str(row.get("Critical", "")).strip().upper()
                    if crit_flag in ["Y", "YES", "TRUE", "1", "HIGH", "LOW"]:
                        is_critical_row = True
                            
                for h in headers:
                    val = str(row.get(h, "")).replace("\n", " ").replace("|", "\\|")
                    if is_critical_row and h == "Result":
                        val = f"<span style='color:red; font-weight:bold;'>{val}</span>"
                    row_vals.append(val)
                lines.append("| " + " | ".join(row_vals) + " |")
            lines.append("")
            
        return "\n".join(lines)

    def _add_kv_paragraph(self, doc, kv_dict: dict):
        """Adds a key-value paragraph block to a docx document."""
        p = doc.add_paragraph()
        for k, v in kv_dict.items():
            val_str = str(v).replace("\n", " ").strip()
            if len(val_str) > 500:
                val_str = val_str[:500] + "..."
            run_key = p.add_run(f"{k}: ")
            run_key.bold = True
            p.add_run(f"{val_str}\n")

    def render_docx(self, data: dict) -> str:
        """Renders the record data into a Word document and returns the temp file path."""
        doc = docx.Document()
        
        title = doc.add_heading('PATIENT MEDICAL RECORD', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_heading('Consolidated Clinical Summary', level=1)
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        
        # --- Demographics ---
        patient_name = str(data['patient'].get('name', '')).upper()
        doc.add_heading(f'1. Patient: {patient_name}', level=2)
        doc.add_heading('Demographics', level=3)
        
        demo = data["patient"].get("demographics", {})
        if demo:
            self._add_kv_paragraph(doc, demo)
        else:
            p = doc.add_paragraph()
            p.add_run('Registration No: ').bold = True
            p.add_run(f"{data['patient']['registration_no']}\n")
            p.add_run('Patient ID: ').bold = True
            p.add_run(f"{data['patient']['patient_id']}\n")
        
        # --- Next of Kin ---
        nok = data["patient"].get("next_of_kin", {})
        if nok:
            doc.add_heading('Next of Kin', level=3)
            self._add_kv_paragraph(doc, nok)
        
        # --- Visit Info ---
        visits = data.get("visits", [])
        section_num = 2
        if visits:
            doc.add_heading(f'{section_num}. Visit History ({len(visits)} visit{"s" if len(visits) > 1 else ""})', level=2)
            for i, v_info in enumerate(visits, 1):
                visit_no = v_info.get("Visit No", f"Visit {i}")
                doc.add_heading(f'Visit: {visit_no}', level=3)
                self._add_kv_paragraph(doc, v_info)
            section_num += 1
        
        # --- Clinical Assessment ---
        assessments = data.get("clinical_assessments", [])
        if assessments:
            doc.add_heading(f'{section_num}. Clinical Assessment ({len(assessments)} record{"s" if len(assessments) > 1 else ""})', level=2)
            for i, asmt in enumerate(assessments, 1):
                if len(assessments) > 1:
                    doc.add_heading(f'Assessment {i}', level=3)
                self._add_kv_paragraph(doc, asmt)
            section_num += 1
        
        # --- Tabular Sections ---
        for sec in data["sections"]:
            row_count = len(sec['rows'])
            if row_count > 0:
                doc.add_heading(f"{section_num}. {sec['title']} ({row_count} readings/records)", level=2)
            else:
                doc.add_heading(f"{section_num}. {sec['title']}", level=2)
            section_num += 1
                
            if sec["error"]:
                doc.add_paragraph(sec["error"])
                continue
                
            if sec.get("format") == "paragraph":
                for row_idx, row in enumerate(sec["rows"]):
                    if len(sec["rows"]) > 1:
                        doc.add_paragraph(f"Record {row_idx + 1}", style='Heading 4')
                    for k, v in row.items():
                        p = doc.add_paragraph()
                        p.add_run(f"{str(k).replace('_', ' ').title()}: ").bold = True
                        p.add_run(str(v))
                continue

            headers = []
            for r in sec["rows"]:
                for k in r.keys():
                    if k not in headers:
                        headers.append(k)
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = 'Table Grid'
            hdr_cells = t.rows[0].cells
            for i, hdr in enumerate(headers):
                hdr_cells[i].text = str(hdr).replace("_", " ").title()
                
            for row in sec["rows"]:
                row_cells = t.add_row().cells
                for i, key in enumerate(headers):
                    row_cells[i].text = str(row.get(key, ""))
                    
            doc.add_paragraph("\n")

        fd, temp_path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc.save(temp_path)
        
        return temp_path

    def render_transcript_docx(self, session) -> str:
        """Renders only the current session chat transcript into a Word document."""
        doc = docx.Document()
        
        title = doc.add_heading('CONSULTATION TRANSCRIPT', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        
        if session.patient:
            patient_name = str(session.patient.name).upper()
            doc.add_heading(f'Patient: {patient_name}', level=2)
            p = doc.add_paragraph()
            p.add_run('Registration No: ').bold = True
            p.add_run(f"{session.patient.registration_no}\n")
            p.add_run('Patient ID: ').bold = True
            p.add_run(f"{session.patient.patient_id}\n")
            
        doc.add_heading('Session History', level=2)
        
        if not session.history:
            doc.add_paragraph("No interactions in this session.")
        
        for idx, turn in enumerate(session.history, 1):
            # Question
            p_q = doc.add_paragraph()
            p_q.add_run(f"Q{idx}: ").bold = True
            p_q.add_run(str(turn.get("question", "")))
            
            # Answer (parse markdown)
            raw_answer = str(turn.get("answer", ""))
            lines = raw_answer.split('\n')
            
            i = 0
            while i < len(lines):
                line = lines[i].strip()
                
                # Table detection: starts with |, next line is |---|
                if line.startswith('|') and len(line) > 1 and i + 1 < len(lines) and lines[i+1].strip().startswith('|---'):
                    headers = [cell.strip().replace('\\|', '|') for cell in line.split('|')[1:-1]]
                    i += 2
                    rows = []
                    while i < len(lines) and lines[i].strip().startswith('|'):
                        rows.append([cell.strip().replace('\\|', '|') for cell in lines[i].split('|')[1:-1]])
                        i += 1
                        
                    if headers:
                        t = doc.add_table(rows=1, cols=len(headers))
                        t.style = 'Table Grid'
                        for j, h in enumerate(headers):
                            # In case headers length mismatch
                            if j < len(t.rows[0].cells):
                                t.rows[0].cells[j].text = h
                        for row in rows:
                            # Pad or trim row to match headers
                            padded_row = row + [""] * max(0, len(headers) - len(row))
                            padded_row = padded_row[:len(headers)]
                            row_cells = t.add_row().cells
                            for j, cell in enumerate(padded_row):
                                row_cells[j].text = cell
                    doc.add_paragraph("") # Space after table
                    continue
                    
                # Bold detection
                elif line:
                    p_a = doc.add_paragraph()
                    if line.startswith('**') and line.endswith('**'):
                        p_a.add_run(line[2:-2]).bold = True
                    else:
                        p_a.add_run(line)
                        
                i += 1
                
            doc.add_paragraph("-" * 40) # separator
            
        fd, temp_path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc.save(temp_path)
        
        return temp_path